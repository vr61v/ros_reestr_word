import openpyxl
import docx


def findRequestNumber(excel):
    return int(excel['Реестр'][f'A{excel["Реестр"].max_row}'].internal_value.split('/')[0]) + 1


def appendInDocx(table, numberRequest, iterator, region, locality, address, date):
    cells = table.rows[iterator - 1].cells
    cells[0].text = f"{iterator}"
    cells[1].text = f"{region}, {locality}, {address}(Запрос №{numberRequest} от {date}"
    cells[2].text = "1- ЕГРН"
    cells[3].text = "500"


top = int(input("First row = "))
bottom = int(input("Last row = "))
count = bottom - top

excel = openpyxl.load_workbook(r"excel.xlsx")

doc = docx.Document()
table = doc.add_table(rows=count, cols=4)
table.style = 'Table Grid'

it = 0
for i in list(excel['Реестр'].rows)[top:bottom]:
    try:
        if i[4].internal_value is not None and "Выписка из ЕГРН" in i[4].internal_value:
            numberRequest = i[0].value
            it += 1
            region = i[1].value
            locality = i[2].value
            address = i[3].value
            date = i[6].value.date()

            appendInDocx(table, numberRequest, it, region, locality, address, date)
    except:
        continue

doc.save(r"docx.docx")
